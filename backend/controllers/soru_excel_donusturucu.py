"""Toplu soru yüklemenin Excel (.xlsx) BİÇİM dönüştürücüsü — Controller yardımcısı.

Neden: Excel bir taşıma biçimidir, iş kuralı değildir. Bu dosya yalnızca iki yönlü
biçim dönüşümü yapar: (1) yüklenen .xlsx baytlarını Service'in beklediği satır
sözlüklerine çevirir, (2) Service'in verdiği dropdown içeriğinden şablon .xlsx
baytlarını üretir. İŞ KURALI, yetki ve sanitizasyon BURADA YOKTUR; hepsi
soru_yukleme_service/soru_service'tedir.

Kullanıcının gördüğü Türkçe sütun başlıkları ile teknik anahtarların eşlemesi TEK
YERDE (bu dosyada) tanımlıdır; şablon üretimi ve okuma aynı kaynağı kullanır,
böylece ikisi ayrışamaz.

Hata yönetimi: Bozuk/uyumsuz dosya kullanıcı hatasıdır -> ValidationError fırlatılır
(loglanmaz, sınırda bir kez loglanır). openpyxl'in ham istisnası kullanıcıya
SIZDIRILMAZ; orijinali yalnızca `from` ile zincirlenir.
"""

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common.constants import MAKS_EXCEL_SORU_SATIRI
from common.errors import ValidationError

# Şıkların yazılacağı sütun sayısı (şablonda "Seçenek 1..15"). Tipe göre kaç şık
# gerektiği bir İŞ KURALIDIR ve Service'te doğrulanır; burada yalnız sütun sayısı.
_SECENEK_SUTUN_SAYISI = 15

# Şık dışındaki sabit sütunların başlığı -> teknik anahtar eşlemesi (sırayla).
_SABIT_SUTUNLAR = (
    ("Soru Tipi", "soru_tipi"),
    ("Konu", "konu"),
    ("Amaç", "amac"),
    ("Soru Metni", "soru_metni"),
)

# Şablonun 1. satırındaki tam başlık dizisi (kullanıcı bunları görür ve okuma
# sırasında birebir bu dizi beklenir).
_BASLIKLAR = [baslik for baslik, _ in _SABIT_SUTUNLAR] + [
    f"Seçenek {sira}" for sira in range(1, _SECENEK_SUTUN_SAYISI + 1)
]

# Sayfa adları: veri sayfası ilk sıradadır (okuma buradan yapılır).
_VERI_SAYFASI = "Sorular"
_YONERGE_SAYFASI = "Yönerge"
_LISTE_SAYFASI = "Listeler"

# Şablondaki "Listeler" sayfasında her dropdown kaynağının sütunu ve başlığı.
# Değerler uzun olabildiğinden dropdown'lar inline liste yerine bu aralıklara
# bağlanır (inline formula1 255 karakter sınırına takılır).
_LISTE_SUTUNLARI = (
    ("soru_tipi_etiketleri", "Soru Tipi", 1),
    ("konular", "Konu", 2),
    ("amaclar", "Amaç", 3),
)

# Şablonun "Yönerge" sayfasındaki açıklama satırları (tipe göre şık kuralları).
_YONERGE_SATIRLARI = (
    "Anket Sorusu Toplu Yükleme Yönergesi",
    "",
    "1) Her satır TEK bir soruyu tanımlar. Başlık satırını (1. satır) silmeyin, "
    "değiştirmeyin.",
    "2) Soru Tipi, Konu ve Amaç hücrelerinde yalnızca açılır listedeki değerleri "
    "kullanın.",
    "3) Soru Metni DÜZ METİNDİR; biçimlendirme (kalın, renk) kaydedilmez. "
    "Hücre içinde satır sonu kullanabilirsiniz.",
    "",
    "Soru tipine göre şık kuralları:",
    "- Çoktan seçmeli (tek/çoklu yanıt) ve Listeden seçmeli soru: Seçenek 1'den "
    "başlayarak 2 ile 15 arasında şık yazın.",
    "- Evet-hayır sorusu: şık sütunlarını BOŞ bırakın; Evet/Hayır sistem "
    "tarafından eklenir.",
    "- 5'li skala sorusu: yalnızca Seçenek 1 ve Seçenek 2'yi doldurun; bunlar 1 ve "
    "5 uçlarının ifadeleridir (2-3-4 sistem tarafından eklenir).",
    "- Yorum sorusu: şık YAZILMAZ; tüm şık sütunları boş kalmalıdır.",
    "- Grid Sorusu Excel ile yüklenemez; listede yer almaz.",
    "",
    "Yükleme ya tamamen başarılı olur ya da hiç kayıt eklenmez: hatalı satır "
    "varsa hiçbir soru yüklenmez ve satır bazlı hata listesi gösterilir.",
    f"Tek dosyada en fazla {MAKS_EXCEL_SORU_SATIRI} soru satırı yüklenebilir.",
)


def _hucreyi_metne_cevir(deger) -> str:
    """Bir Excel hücre değerini kenar boşlukları alınmış metne çevirir (None -> '')."""
    if deger is None:
        return ""
    return str(deger).strip()


def _baslik_satirini_dogrula(satir) -> None:
    """Şablonun başlık satırının beklenen sütunlarla birebir aynı olduğunu doğrular.

    Neden: Sütunlar taşınmış/silinmişse veriler yanlış alanlara eşlenir; bu sessiz
    bozulma yerine açık bir hata verilir. Sondaki tamamen boş hücreler yok sayılır
    (Excel dosyaları çoğu zaman kuyrukta boş sütun taşır).
    """
    hucreler = [_hucreyi_metne_cevir(deger) for deger in satir or ()]
    while hucreler and not hucreler[-1]:
        hucreler.pop()
    if hucreler != _BASLIKLAR:
        raise ValidationError(
            "Şablon sütunları değiştirilmiş. Lütfen indirdiğiniz şablonu "
            "kullanın ve başlık satırını değiştirmeyin."
        )


def _veri_satirini_cevir(satir, satir_no: int) -> dict | None:
    """Bir veri satırını Service'in beklediği sözlüğe çevirir; tamamen boşsa None döner.

    Şık sütunlarından BOŞ olanlar listeye konmaz; dolu olanlar soldan sağa sırayla
    alınır. satir_no kullanıcının Excel'de gördüğü GERÇEK satır numarasıdır, böylece
    hata raporu doğrudan dosyayla eşleşir.
    """
    hucreler = [_hucreyi_metne_cevir(deger) for deger in satir or ()]
    if not any(hucreler):
        return None

    # Eksik hücreler (kısa satır) boş metinle tamamlanır ki sütun eşlemesi kaymasın.
    hucreler += [""] * (len(_BASLIKLAR) - len(hucreler))

    kayit = {"satir_no": satir_no}
    for sutun_indeksi, (_, anahtar) in enumerate(_SABIT_SUTUNLAR):
        kayit[anahtar] = hucreler[sutun_indeksi]

    ilk_secenek_indeksi = len(_SABIT_SUTUNLAR)
    kayit["secenekler"] = [
        metin
        for metin in hucreler[
            ilk_secenek_indeksi : ilk_secenek_indeksi + _SECENEK_SUTUN_SAYISI
        ]
        if metin
    ]
    return kayit


def oku_soru_satirlari(dosya_baytlari: bytes) -> list[dict]:
    """Yüklenen .xlsx baytlarını soru satırı sözlüklerine çevirir (yalnız BİÇİM).

    İlk sayfa veri sayfası kabul edilir; 1. satır başlık olarak doğrulanır. Tamamen
    boş satırlar atlanır (Excel'in kuyruk satırları). Veri satırı sayısı
    MAKS_EXCEL_SORU_SATIRI'yı aşarsa dosya doğrulanmadan reddedilir. Değerlerin
    ANLAMI (geçerli tip, zorunlu alan, şık adedi) burada DEĞİL Service'te
    doğrulanır. Bozuk/açılamayan dosya ValidationError'a çevrilir; openpyxl'in ham
    istisnası kullanıcıya sızmaz (yalnızca `from` ile zincirlenir).
    """
    try:
        calisma_kitabi = load_workbook(
            BytesIO(dosya_baytlari), read_only=True, data_only=True
        )
    except Exception as hata:  # noqa: BLE001 - her bozuk dosya kullanıcı hatasıdır
        raise ValidationError(
            "Excel dosyası okunamadı. Lütfen geçerli bir .xlsx dosyası yükleyin."
        ) from hata

    try:
        sayfa = calisma_kitabi.worksheets[0]
        satirlar = sayfa.iter_rows(values_only=True)
        _baslik_satirini_dogrula(next(satirlar, None))

        kayitlar: list[dict] = []
        # Başlık 1. satır olduğundan veri satırları 2'den başlar.
        for satir_no, satir in enumerate(satirlar, start=2):
            kayit = _veri_satirini_cevir(satir, satir_no)
            if kayit is None:
                continue
            kayitlar.append(kayit)
            if len(kayitlar) > MAKS_EXCEL_SORU_SATIRI:
                raise ValidationError(
                    f"Tek seferde en fazla {MAKS_EXCEL_SORU_SATIRI} soru "
                    "yüklenebilir."
                )
    finally:
        calisma_kitabi.close()

    return kayitlar


def _yaz_yonerge_sayfasi(calisma_kitabi: Workbook) -> None:
    """Şablona, tipe göre şık kurallarını anlatan "Yönerge" sayfasını ekler."""
    sayfa = calisma_kitabi.create_sheet(_YONERGE_SAYFASI)
    for metin in _YONERGE_SATIRLARI:
        sayfa.append([metin])
    sayfa.column_dimensions["A"].width = 110


def _yaz_liste_sayfasi(calisma_kitabi: Workbook, sablon_verisi: dict) -> dict:
    """Dropdown kaynak değerlerini "Listeler" sayfasına yazar; aralık referanslarını döner.

    Neden ayrı sayfa: DataValidation'ın inline listesi (formula1) 255 karakterle
    sınırlıdır; konu/amaç değerleri bu sınırı kolayca aşar. Değerler sayfaya yazılıp
    dropdown'lar aralığa bağlanır. Dönen sözlük: veri anahtarı -> formül referansı
    (değer yoksa o anahtar hiç bulunmaz, dropdown da eklenmez).
    """
    sayfa = calisma_kitabi.create_sheet(_LISTE_SAYFASI)
    referanslar: dict[str, str] = {}

    for anahtar, baslik, sutun_no in _LISTE_SUTUNLARI:
        sutun_harfi = get_column_letter(sutun_no)
        sayfa.cell(row=1, column=sutun_no, value=baslik)
        degerler = sablon_verisi.get(anahtar) or []
        for sira, deger in enumerate(degerler, start=2):
            sayfa.cell(row=sira, column=sutun_no, value=deger)
        if degerler:
            son_satir = len(degerler) + 1
            referanslar[anahtar] = (
                f"={_LISTE_SAYFASI}!${sutun_harfi}$2:${sutun_harfi}${son_satir}"
            )
        sayfa.column_dimensions[sutun_harfi].width = 40

    return referanslar


def _ekle_dropdownlar(sayfa, referanslar: dict) -> None:
    """Veri sayfasının Soru Tipi/Konu/Amaç sütunlarına açılır liste kısıtı ekler.

    Kısıt, kullanıcıya kolaylıktır; asıl doğrulama sunucuda (Service) yapılır.
    Değeri olmayan kategori için dropdown eklenmez (boş liste Excel'i bozar).
    """
    son_satir = MAKS_EXCEL_SORU_SATIRI + 1
    for anahtar, _, sutun_no in _LISTE_SUTUNLARI:
        formul = referanslar.get(anahtar)
        if not formul:
            continue
        dogrulama = DataValidation(type="list", formula1=formul, allow_blank=True)
        sayfa.add_data_validation(dogrulama)
        sutun_harfi = get_column_letter(sutun_no)
        dogrulama.add(f"{sutun_harfi}2:{sutun_harfi}{son_satir}")


def uret_sablon_baytlari(sablon_verisi: dict) -> bytes:
    """Service'in verdiği dropdown içeriğinden şablon .xlsx baytlarını üretir.

    Üç sayfa: "Sorular" (başlıklar + dropdown'lı boş veri alanı), "Yönerge" (tipe
    göre şık kuralları, soru metninin düz metin olduğu) ve "Listeler" (dropdown
    kaynak değerleri). İçeriğin ANLAMI Service'ten gelir; burada yalnızca Excel
    yerleşimi kurulur.
    """
    calisma_kitabi = Workbook()
    veri_sayfasi = calisma_kitabi.active
    veri_sayfasi.title = _VERI_SAYFASI
    veri_sayfasi.append(_BASLIKLAR)
    for sutun_no in range(1, len(_BASLIKLAR) + 1):
        veri_sayfasi.column_dimensions[get_column_letter(sutun_no)].width = 28

    _yaz_yonerge_sayfasi(calisma_kitabi)
    referanslar = _yaz_liste_sayfasi(calisma_kitabi, sablon_verisi)
    _ekle_dropdownlar(veri_sayfasi, referanslar)

    tampon = BytesIO()
    calisma_kitabi.save(tampon)
    return tampon.getvalue()
